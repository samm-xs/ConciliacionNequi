import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from tabulate import tabulate


def limpiar_dinero(valor):
    if valor is None:
        return None

    valor = str(valor).strip()

    if valor == "":
        return None

    valor = valor.replace("$", "").replace(",", "").strip()

    try:
        return float(valor)
    except ValueError:
        return None


def extraer_movimientos_nequi(pdf_path, password=None):
    tablas_movimientos = []

    with pdfplumber.open(pdf_path, password=password) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                if not table:
                    continue

                header = [str(x).strip() if x else "" for x in table[0]]

                if "Fecha del movimiento" in header and "Descripción" in header:
                    tablas_movimientos.append(table)

    dfs = [pd.DataFrame(table[1:], columns=table[0]) for table in tablas_movimientos]
    df = pd.concat(dfs, ignore_index=True)
    df.columns = [str(col).strip() for col in df.columns]
    df.insert(0, "id", range(1, len(df) + 1))
    df["Fecha del movimiento"] = pd.to_datetime(
        df["Fecha del movimiento"],
        format="%d/%m/%Y",
        errors="coerce",
    )

    for col in ["Valor", "Saldo"]:
        if col in df.columns:
            df[col] = df[col].apply(limpiar_dinero)

    return df


def exportar_excel(df, nombre_archivo="Extracto_nequiMarzo.xlsx"):
    df_excel = df.copy()

    if "Fecha del movimiento" in df_excel.columns:
        df_excel["Fecha del movimiento"] = df_excel["Fecha del movimiento"].dt.date

    with pd.ExcelWriter(nombre_archivo, engine="openpyxl") as writer:
        df_excel.to_excel(writer, index=False, sheet_name="Movimientos")
        ws = writer.sheets["Movimientos"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        header_font = Font(bold=True, color="000000")
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")

        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value == "Fecha del movimiento":
                col_letter = get_column_letter(col_idx)
                for c in ws[col_letter][1:]:
                    c.number_format = "yyyy-mm-dd"

        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value in ["Valor", "Saldo"]:
                col_letter = get_column_letter(col_idx)
                for c in ws[col_letter][1:]:
                    c.number_format = "#,##0.00"

        for column_cells in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                value = cell.value

                if value is None:
                    continue

                max_length = max(max_length, len(str(value)))

            ws.column_dimensions[column_letter].width = max_length + 3

    print(f"Archivo exportado: {nombre_archivo}")
