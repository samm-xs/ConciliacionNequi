import pdfplumber
import pandas as pd
import re
from tabulate import tabulate
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


MESES = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "SETIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}


PATRON_MOVIMIENTO = re.compile(
    r"^"
    r"(?P<dia>\d{2})\s+"
    r"(?P<descripcion>.+?)\s+"
    r"(?P<oficina>INTERNET BANCA|ORITO)\s+"
    r"(?P<no_dcto>\d+)\s+"
    r"(?P<valor>-?[\d.]+,\d{2})\s+"
    r"(?P<saldo>-?[\d.]+,\d{2})"
    r"$"
)


def detectar_mes_y_anio(pdf):
    texto = ""

    for page in pdf.pages[:3]:
        texto += "\n" + (page.extract_text() or "")

    texto = texto.upper()

    for nombre_mes, numero_mes in MESES.items():
        patron = rf"{nombre_mes}\s+(\d{{4}})"
        match = re.search(patron, texto)

        if match:
            return numero_mes, int(match.group(1))

    raise ValueError("No pude detectar el mes y año del extracto.")


def limpiar_valor(valor):
    """
    Convierte:
    152.200,00     -> 152200.00
    -4.152.508,00  -> -4152508.00
    7.416.422,37   -> 7416422.37
    """

    if valor is None:
        return None

    valor = str(valor).strip()

    if valor == "":
        return None

    valor = re.sub(r"[^\d,.-]", "", valor)

    if valor == "":
        return None

    valor = valor.replace(".", "").replace(",", ".")

    try:
        return float(valor)
    except ValueError:
        return None


def extraer_extracto_agrario(pdf_path):
    filas = []

    with pdfplumber.open(pdf_path) as pdf:
        mes, anio = detectar_mes_y_anio(pdf)

        for page in pdf.pages:
            texto = page.extract_text() or ""

            for linea in texto.splitlines():
                linea = " ".join(linea.split()).strip()

                match = PATRON_MOVIMIENTO.match(linea)

                if not match:
                    continue

                datos = match.groupdict()

                dia = int(datos["dia"])

                fila = {
                    "fecha": pd.Timestamp(year=anio, month=mes, day=dia),
                    "descripcion": datos["descripcion"].strip(),
                    "oficina": datos["oficina"].strip(),
                    "no_dcto": datos["no_dcto"].strip(),
                    "valor": limpiar_valor(datos["valor"]),
                    "saldo": limpiar_valor(datos["saldo"]),
                }

                filas.append(fila)

    df = pd.DataFrame(filas)

    df.insert(0, "id", range(1, len(df) + 1))

    df = df[
        [
            "id",
            "fecha",
            "descripcion",
            "oficina",
            "no_dcto",
            "valor",
            "saldo",
        ]
    ]

    return df


def exportar_excel(df, nombre_archivo="Extracto_bancoAgrarioDiciembre.xlsx"):
    df_excel = df.copy()

    df_excel["fecha"] = df_excel["fecha"].dt.date

    with pd.ExcelWriter(nombre_archivo, engine="openpyxl") as writer:
        df_excel.to_excel(writer, index=False, sheet_name="Movimientos")

        ws = writer.sheets["Movimientos"]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_fill = PatternFill(
            start_color="FFFFFF",
            end_color="FFFFFF",
            fill_type="solid",
        )

        header_font = Font(bold=True, color="000000")

        thin = Side(style="thin", color="D9D9D9")

        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin,
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")

        for col_idx, cell in enumerate(ws[1], start=1):
            col_letter = get_column_letter(col_idx)

            if cell.value == "fecha":
                for c in ws[col_letter][1:]:
                    c.number_format = "yyyy-mm-dd"

            if cell.value in ["valor", "saldo"]:
                for c in ws[col_letter][1:]:
                    c.number_format = "#,##0.00"

        anchos = {
            "A": 8,
            "B": 14,
            "C": 50,
            "D": 18,
            "E": 18,
            "F": 14,
            "G": 16,
            "H": 16,
        }

        for col, ancho in anchos.items():
            ws.column_dimensions[col].width = ancho

    print(f"Archivo exportado: {nombre_archivo}")

if __name__ == "__main__":
    df = extraer_extracto_agrario("C:/Users/Sebastian/Downloads/Cuenta Corriente Diciembre 2025.pdf")
    exportar_excel(df)