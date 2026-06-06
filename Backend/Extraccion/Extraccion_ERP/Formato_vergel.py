import pdfplumber
import pandas as pd
import re
from tabulate import tabulate
from openpyxl.utils import get_column_letter


pdf_path = "C:/Users/Sebastian/Downloads/Marzo sistema.PDF"


COLUMNAS = {
    "fecha": (20, 70),
    "doc": (70, 95),
    "numero": (95, 135),
    "tercero": (135, 195),
    "nombre_tercero": (195, 305),
    "doc_refer": (305, 350),
    "detalle": (350, 455),
    "debitos": (455, 520),
    "creditos": (520, 595),
}


PATRON_FECHA = re.compile(r"^\d{2}/\d{2}/\d{4}$")


def agrupar_por_linea(words, tolerancia=3):
    lineas = []

    for palabra in sorted(words, key=lambda w: (w["top"], w["x0"])):
        agregada = False

        for linea in lineas:
            if abs(linea[0]["top"] - palabra["top"]) <= tolerancia:
                linea.append(palabra)
                agregada = True
                break

        if not agregada:
            lineas.append([palabra])

    return [sorted(linea, key=lambda w: w["x0"]) for linea in lineas]


def limpiar_valor(valor):
    """
    Convierte:
    96.300.00 -> 96300.00
    1.952.795.00 -> 1952795.00
    139.269.47 -> 139269.47
    """

    if valor is None:
        return None

    valor = str(valor).strip()

    if valor == "":
        return None

    valor = re.sub(r"[^\d.,-]", "", valor)

    if valor == "":
        return None

    partes = valor.split(".")

    if len(partes) > 2:
        valor_limpio = "".join(partes[:-1]) + "." + partes[-1]
        return float(valor_limpio)

    try:
        return float(valor)
    except ValueError:
        return None


def extraer_fila(linea):
    if not linea:
        return None

    primer_texto = linea[0]["text"]

    if not PATRON_FECHA.match(primer_texto):
        return None

    fila = {
        "fecha": [],
        "doc": [],
        "numero": [],
        "tercero": [],
        "nombre_tercero": [],
        "doc_refer": [],
        "detalle": [],
        "debitos": [],
        "creditos": [],
    }

    for palabra in linea:
        x = palabra["x0"]
        texto = palabra["text"]

        for columna, (inicio, fin) in COLUMNAS.items():
            if inicio <= x < fin:
                fila[columna].append(texto)
                break

    fila = {
        columna: " ".join(partes).strip()
        for columna, partes in fila.items()
    }

    if not fila["fecha"] or not fila["doc"] or not fila["tercero"]:
        return None

    return fila


def extraer_libro_auxiliar(pdf_path):
    filas = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            words = page.extract_words(
                x_tolerance=1,
                y_tolerance=3,
                keep_blank_chars=False,
                use_text_flow=False,
            )

            lineas = agrupar_por_linea(words)

            for linea in lineas:
                fila = extraer_fila(linea)

                if fila is not None:
                    filas.append(fila)

    df = pd.DataFrame(filas)

    df.insert(0, "id", range(1, len(df) + 1))

    df["fecha"] = pd.to_datetime(
        df["fecha"],
        format="%d/%m/%Y",
        errors="coerce"
    )

    df["debitos"] = df["debitos"].apply(limpiar_valor)
    df["creditos"] = df["creditos"].apply(limpiar_valor)

    df = df[
        [
            "id",
            "fecha",
            "doc",
            "numero",
            "tercero",
            "nombre_tercero",
            "doc_refer",
            "detalle",
            "debitos",
            "creditos",
        ]
    ]

    return df


# def preparar_vista(df):
#     df_vista = df.copy()

#     df_vista["fecha"] = df_vista["fecha"].dt.strftime("%Y-%m-%d")

#     for columna in ["debitos", "creditos"]:
#         df_vista[columna] = df_vista[columna].apply(
#             lambda x: "" if pd.isna(x) else f"{x:,.2f}"
#         )

#     return df_vista


# def mostrar_tabla(df, limite=30):
#     df_vista = preparar_vista(df)

#     if limite is not None:
#         df_mostrar = df_vista.head(limite)
#     else:
#         df_mostrar = df_vista

#     print(
#         tabulate(
#             df_mostrar,
#             headers="keys",
#             tablefmt="grid",
#             showindex=False
#         )
#     )

#     if limite is not None and limite < len(df):
#         print(f"[mostrando {limite} de {len(df)} registros x {len(df.columns)} columnas]")
#     else:
#         print(f"[{len(df)} registros x {len(df.columns)} columnas]")


# def exportar_excel(df, nombre_archivo="libro_auxiliarNequiMarzo.xlsx"):
#     df_excel = df.copy()

#     # Para que Excel no muestre 2025-08-01 00:00:00
#     df_excel["fecha"] = df_excel["fecha"].dt.date

#     with pd.ExcelWriter(nombre_archivo, engine="openpyxl") as writer:
#         df_excel.to_excel(writer, index=False, sheet_name="Libro Auxiliar")

#         ws = writer.sheets["Libro Auxiliar"]

#         # Congelar encabezado
#         ws.freeze_panes = "A2"

#         # Auto filtro
#         ws.auto_filter.ref = ws.dimensions

#         # Anchos de columna
#         anchos = {
#             "A": 8,
#             "B": 14,
#             "C": 10,
#             "D": 12,
#             "E": 16,
#             "F": 32,
#             "G": 14,
#             "H": 22,
#             "I": 16,
#             "J": 16,
#         }

#         for col, ancho in anchos.items():
#             ws.column_dimensions[col].width = ancho

#         # Formato de fecha
#         for celda in ws["B"][1:]:
#             celda.number_format = "yyyy-mm-dd"

#         # Formato de dinero con dos decimales
#         for col in ["I", "J"]:
#             for celda in ws[col][1:]:
#                 celda.number_format = "#,##0.00"

#     print(f"Archivo exportado: {nombre_archivo}")


# df = extraer_libro_auxiliar(pdf_path)

# mostrar_tabla(df, limite=30)

##df.to_json("libro_auxiliarNequiMarzo.json", orient="records", date_format="iso", force_ascii=False)

##exportar_excel(df)