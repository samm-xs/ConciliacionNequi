from pathlib import Path

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


COLOR_ENCABEZADO = "1F4E78"
COLOR_CONCILIADO = "D9EAD3"
COLOR_POR_REVISAR = "FFF2CC"
COLOR_ERROR = "F4CCCC"
COLOR_OK = "D9EAD3"
COLOR_PENDIENTE = "E7E6E6"
COLOR_BORDE = "000000"
COLOR_TITULO = "A64D79"
COLOR_BLOQUE_LIBRO = "3C78D8"
COLOR_BLOQUE_AUTOMATICO = "A64D79"
COLOR_BLOQUE_MANUAL = "BF9000"
MAX_COMPUESTOS_EXPORTABLES = 1000
TAMANO_FUENTE_DATOS = 13
TAMANO_FUENTE_ENCABEZADO = 14
TAMANO_FUENTE_TITULO = 20

HOJA_LIBRO_AUXILIAR_BANCO = "Libro auxiliar vs banco"
HOJA_CONCEPTOS_BANCARIOS = "Conceptos bancarios"
HOJA_CONCILIACION_COMPUESTOS = "Conciliacion de compuestos"
HOJA_REGISTROS_PENDIENTES_BANCO = "Registros pendientes del banco"
HOJA_EXTRACTOS_BANCO = "Extractos banco"
HOJA_LIBRO_AUXILIAR_EXTRAIDO = "Libro Auxiliar"

COLUMNAS_LIBRO_AUXILIAR = [
    "ID",
    "Fecha",
    "Nombre",
    "Valor",
    "Estado",
    "Conciliado del banco con registros:",
    "Estado",
    "Conciliado del banco con registros:",
    "Estado automatico base",
]

ANCHOS_LIBRO_AUXILIAR = {
    "A": 10,
    "B": 13,
    "C": 28,
    "D": 16,
    "E": 14,
    "F": 65,
    "G": 16,
    "H": 55,
}

FILA_TITULO_LIBRO_AUXILIAR = 1
FILA_BLOQUES_LIBRO_AUXILIAR = 2
FILA_ENCABEZADOS_LIBRO_AUXILIAR = 3
FILA_DATOS_LIBRO_AUXILIAR = 4
COL_ESTADO_AUTOMATICO = 5
COL_CONCILIADO_AUTOMATICO = 6
COL_ESTADO_MANUAL = 7
COL_CONCILIADO_MANUAL = 8
COL_ESTADO_AUTOMATICO_BASE = 9

COLUMNAS_CONCEPTOS_BANCARIOS = [
    "Concepto",
    "Cantidad movimientos",
    "Valor total",
    "Descripcion",
]

COLUMNAS_REGISTROS_PENDIENTES_BANCO = [
    "ID",
    "Fecha",
    "Nombre",
    "Valor",
    "Estado",
]

ANCHOS_REGISTROS_PENDIENTES_BANCO = {
    "A": 10,
    "B": 13,
    "C": 32,
    "D": 14,
    "E": 36,
}

FILA_TITULO_REGISTROS_BANCO = 1
FILA_ENCABEZADOS_REGISTROS_BANCO = 2
FILA_DATOS_REGISTROS_BANCO = 3

COLUMNAS_COMPUESTOS_VISIBLES = [
    "Movimientos banco",
    "Movimientos libro_aux",
    "Diferencia",
    "Estado",
    "Validacion",
    "Descripcion",
    "Observacion",
]

COLUMNAS_COMPUESTOS_AUXILIARES = [
    "ids banco",
    "ids libro_aux",
    "id banco 1",
    "id banco 2",
    "id banco 3",
    "id libro_aux 1",
    "id libro_aux 2",
    "id libro_aux 3",
]

COLUMNAS_COMPUESTOS = COLUMNAS_COMPUESTOS_VISIBLES + COLUMNAS_COMPUESTOS_AUXILIARES

ANCHOS_CONCILIACION_COMPUESTOS = {
    "A": 55,
    "B": 55,
    "C": 14,
    "D": 16,
    "E": 24,
    "F": 42,
    "G": 32,
    "H": 14,
    "I": 14,
    "J": 12,
    "K": 12,
    "L": 12,
    "M": 12,
    "N": 12,
    "O": 12,
}

DESCRIPCION_COMPUESTO_MISMA_FECHA = (
    "Posible compuesto detectado por suma exacta en la misma fecha. "
    "Revisar antes de conciliar."
)


def formatear_fecha(valor):
    fecha = pd.to_datetime(valor, errors="coerce")

    if pd.isna(fecha):
        return ""

    return fecha.date()


def formatear_id(valor):
    if pd.isna(valor):
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return int(valor)

    return valor


def formatear_valor_excel(valor):
    if pd.isna(valor):
        return ""

    return float(valor)


def formatear_valor_compacto(valor):
    if pd.isna(valor):
        return ""

    numero = float(valor)

    if numero.is_integer():
        return f"{numero:,.0f}"

    return f"{numero:,.2f}"


def normalizar_clave_id(valor):
    if pd.isna(valor):
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def dividir_ids(valor, cantidad=3):
    if pd.isna(valor):
        partes = []
    else:
        partes = [parte.strip() for parte in str(valor).split(",") if parte.strip()]

    ids = []
    for parte in partes[:cantidad]:
        try:
            numero = float(parte)
            ids.append(int(numero) if numero.is_integer() else numero)
        except ValueError:
            ids.append(formatear_id(parte))

    while len(ids) < cantidad:
        ids.append("")

    return ids


def construir_lookup_movimientos(df_resultado, prefijo):
    id_columna = f"id_{prefijo}"
    fecha_columna = f"fecha_{prefijo}"
    nombre_columna = f"nombre_{prefijo}"
    valor_columna = f"valor_{prefijo}"
    lookup = {}

    if id_columna not in df_resultado.columns:
        return lookup

    df_movimientos = df_resultado[df_resultado[id_columna].notna()].copy()

    for _, fila in df_movimientos.iterrows():
        clave = normalizar_clave_id(fila[id_columna])

        if clave in lookup:
            continue

        lookup[clave] = {
            "id": formatear_id(fila[id_columna]),
            "fecha": formatear_fecha(fila.get(fecha_columna)),
            "valor": fila.get(valor_columna),
            "nombre": fila.get(nombre_columna, ""),
        }

    return lookup


def formatear_movimiento_compacto(id_movimiento, lookup):
    clave = normalizar_clave_id(id_movimiento)
    movimiento = lookup.get(clave)

    if movimiento is None:
        return str(id_movimiento)

    fecha = movimiento["fecha"]
    fecha_texto = str(fecha) if fecha != "" else ""
    nombre = movimiento["nombre"] if not pd.isna(movimiento["nombre"]) else ""

    return (
        f"{movimiento['id']} | {fecha_texto} | "
        f"{formatear_valor_compacto(movimiento['valor'])} | {nombre}"
    )


def construir_celda_movimientos(ids, lookup):
    return "\n".join(
        formatear_movimiento_compacto(id_movimiento, lookup)
        for id_movimiento in ids
        if id_movimiento != ""
    )


def formatear_descripcion_compuesto(descripcion):
    if not isinstance(descripcion, str):
        return descripcion

    if descripcion.startswith("Posible compuesto:"):
        return DESCRIPCION_COMPUESTO_MISMA_FECHA

    return descripcion


def seleccionar_fila_libro_auxiliar(grupo):
    conciliados = grupo[grupo["estado"].eq("conciliado")]
    if not conciliados.empty:
        return conciliados.iloc[0]

    huerfanos = grupo[grupo["estado"].eq("huerfano_erp")]
    if not huerfanos.empty:
        return huerfanos.iloc[0]

    return grupo.iloc[0]


def preparar_libro_auxiliar_vs_banco(df_resultado):
    df_libro = df_resultado[df_resultado["id_erp"].notna()].copy()

    if df_libro.empty:
        return pd.DataFrame(columns=COLUMNAS_LIBRO_AUXILIAR)

    lookup_banco = construir_lookup_movimientos(df_resultado, "banco")
    filas = []

    for _, grupo in df_libro.groupby(
        df_libro["id_erp"].apply(normalizar_clave_id),
        sort=False,
    ):
        fila = seleccionar_fila_libro_auxiliar(grupo)
        estado = "conciliado" if fila["estado"] == "conciliado" else "por revisar"
        conciliado_con = ""

        if estado == "conciliado" and not pd.isna(fila.get("id_banco")):
            conciliado_con = formatear_movimiento_compacto(
                fila["id_banco"],
                lookup_banco,
            )

        filas.append(
            [
                formatear_id(fila["id_erp"]),
                formatear_fecha(fila.get("fecha_erp")),
                fila.get("nombre_erp", ""),
                formatear_valor_excel(fila.get("valor_erp")),
                estado,
                conciliado_con,
                "por revisar",
                "",
                estado,
            ]
        )

    return pd.DataFrame(filas, columns=COLUMNAS_LIBRO_AUXILIAR)


def preparar_conceptos_bancarios(df_conceptos_banco):
    if df_conceptos_banco is None or df_conceptos_banco.empty:
        return pd.DataFrame(columns=COLUMNAS_CONCEPTOS_BANCARIOS)

    df_excel = pd.DataFrame()
    df_excel["Concepto"] = df_conceptos_banco["concepto"]
    df_excel["Cantidad movimientos"] = df_conceptos_banco["cantidad movimientos"]
    df_excel["Valor total"] = df_conceptos_banco["valor total"]
    df_excel["Descripcion"] = (
        "Concepto bancario informativo. No se concilia contra libro auxiliar."
    )

    return df_excel[COLUMNAS_CONCEPTOS_BANCARIOS]


def preparar_registros_pendientes_banco(df_resultado):
    if df_resultado.empty or "id_banco" not in df_resultado.columns:
        return pd.DataFrame(columns=COLUMNAS_REGISTROS_PENDIENTES_BANCO)

    df_banco = df_resultado[df_resultado["id_banco"].notna()].copy()

    if df_banco.empty:
        return pd.DataFrame(columns=COLUMNAS_REGISTROS_PENDIENTES_BANCO)

    ids_conciliados = set(
        df_banco[df_banco["estado"].eq("conciliado")]["id_banco"].apply(
            normalizar_clave_id
        )
    )
    df_pendientes = df_banco[
        ~df_banco["id_banco"].apply(normalizar_clave_id).isin(ids_conciliados)
    ].copy()

    if df_pendientes.empty:
        return pd.DataFrame(columns=COLUMNAS_REGISTROS_PENDIENTES_BANCO)

    df_pendientes["_clave_id_banco"] = df_pendientes["id_banco"].apply(
        normalizar_clave_id
    )
    df_pendientes = df_pendientes.drop_duplicates(
        subset="_clave_id_banco",
        keep="first",
    )

    df_excel = pd.DataFrame()
    df_excel["ID"] = df_pendientes["id_banco"].apply(formatear_id)
    df_excel["Fecha"] = df_pendientes["fecha_banco"].apply(formatear_fecha)
    df_excel["Nombre"] = df_pendientes["nombre_banco"]
    df_excel["Valor"] = df_pendientes["valor_banco"].apply(formatear_valor_excel)
    df_excel["Estado"] = "disponible"

    return df_excel[COLUMNAS_REGISTROS_PENDIENTES_BANCO]


def preparar_conciliacion_compuestos(df_compuestos, df_resultado):
    if df_compuestos is None or df_compuestos.empty:
        return pd.DataFrame(columns=COLUMNAS_COMPUESTOS)

    if len(df_compuestos) > MAX_COMPUESTOS_EXPORTABLES:
        df_compuestos = df_compuestos.head(MAX_COMPUESTOS_EXPORTABLES).copy()

    lookup_banco = construir_lookup_movimientos(df_resultado, "banco")
    lookup_libro_aux = construir_lookup_movimientos(df_resultado, "erp")
    df_excel = pd.DataFrame()
    df_excel["ids banco"] = df_compuestos["ids_banco"]
    df_excel["ids libro_aux"] = df_compuestos["ids_erp"]
    df_excel["Diferencia"] = df_compuestos["diferencia"].apply(formatear_valor_excel)
    df_excel["Estado"] = "por revisar"
    df_excel["Validacion"] = ""
    df_excel["Descripcion"] = df_compuestos["descripcion"].apply(
        formatear_descripcion_compuesto
    )
    df_excel["Observacion"] = ""

    ids_banco = df_excel["ids banco"].apply(dividir_ids)
    ids_libro_aux = df_excel["ids libro_aux"].apply(dividir_ids)

    df_excel["Movimientos banco"] = ids_banco.apply(
        lambda ids: construir_celda_movimientos(ids, lookup_banco)
    )
    df_excel["Movimientos libro_aux"] = ids_libro_aux.apply(
        lambda ids: construir_celda_movimientos(ids, lookup_libro_aux)
    )

    for indice in range(3):
        df_excel[f"id banco {indice + 1}"] = ids_banco.apply(lambda ids: ids[indice])
        df_excel[f"id libro_aux {indice + 1}"] = ids_libro_aux.apply(
            lambda ids: ids[indice]
        )

    return df_excel[COLUMNAS_COMPUESTOS]


def escribir_dataframe(writer, df, nombre_hoja, startrow=0):
    df.to_excel(writer, index=False, sheet_name=nombre_hoja, startrow=startrow)
    return writer.sheets[nombre_hoja]


def escribir_documento_extraido(writer, df, nombre_hoja):
    if df is None:
        return None

    return escribir_dataframe(writer, df, nombre_hoja)


def ajustar_ancho_columnas(ws):
    for columna in ws.columns:
        letra = get_column_letter(columna[0].column)
        ancho_maximo = 0

        for celda in columna:
            if celda.value is None:
                continue

            ancho_maximo = max(ancho_maximo, len(str(celda.value)))

        ws.column_dimensions[letra].width = min(max(ancho_maximo + 2, 12), 55)


def aplicar_formatos_numericos(ws, fila_encabezados=1):
    for celda in ws[fila_encabezados]:
        encabezado = str(celda.value).lower()
        letra = get_column_letter(celda.column)

        if "fecha" in encabezado:
            for item in ws[letra][1:]:
                item.number_format = "yyyy-mm-dd"

        if encabezado.startswith("valor") or encabezado == "diferencia":
            for item in ws[letra][1:]:
                item.number_format = "#,##0.00"


def aplicar_estilo_base(ws, fila_encabezados=1, fila_datos=2):
    ws.freeze_panes = f"A{fila_datos}"
    ws.auto_filter.ref = (
        f"A{fila_encabezados}:{get_column_letter(ws.max_column)}{ws.max_row}"
    )
    ws.sheet_view.showGridLines = False

    borde_fino = Side(style="thin", color=COLOR_BORDE)
    borde = Border(
        left=borde_fino,
        right=borde_fino,
        top=borde_fino,
        bottom=borde_fino,
    )
    encabezado_fill = PatternFill(
        start_color=COLOR_ENCABEZADO,
        end_color=COLOR_ENCABEZADO,
        fill_type="solid",
    )

    for celda in ws[fila_encabezados]:
        celda.fill = encabezado_fill
        celda.font = Font(
            bold=True,
            size=TAMANO_FUENTE_ENCABEZADO,
            color="FFFFFF",
        )
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    for fila in ws.iter_rows(min_row=fila_datos):
        for celda in fila:
            celda.font = Font(size=TAMANO_FUENTE_DATOS)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    aplicar_formatos_numericos(ws, fila_encabezados=fila_encabezados)
    ajustar_ancho_columnas(ws)


def ocultar_columnas(ws, columnas):
    for nombre_columna in columnas:
        for celda in ws[1]:
            if celda.value == nombre_columna:
                letra = get_column_letter(celda.column)
                ws.column_dimensions[letra].hidden = True
                break


def aplicar_anchos_conciliacion_compuestos(ws):
    for columna, ancho in ANCHOS_CONCILIACION_COMPUESTOS.items():
        ws.column_dimensions[columna].width = ancho


def agregar_validacion_estado_compuesto(ws_compuestos):
    if ws_compuestos.max_row < 2:
        return

    columna_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    rango_destino = f"{columna_estado}2:{columna_estado}{ws_compuestos.max_row}"

    validacion = DataValidation(
        type="list",
        formula1='"por revisar,conciliado"',
        allow_blank=False,
    )
    validacion.error = "Seleccioná conciliado o por revisar."
    validacion.errorTitle = "Estado inválido"
    validacion.prompt = "Marcá conciliado solo si el compuesto fue revisado."
    validacion.promptTitle = "Estado"

    ws_compuestos.add_data_validation(validacion)
    validacion.add(rango_destino)


def crear_formula_id_repetido_en_compuestos(
    fila,
    col_estado,
    columnas_ids,
    mensaje_error,
    ultima_fila,
):
    condiciones = []

    for columna_actual in columnas_ids:
        celda_actual = f"{columna_actual}{fila}"
        conteos = [
            (
                f"COUNTIFS(${col_estado}$2:${col_estado}${ultima_fila},"
                f'"conciliado",'
                f"${columna_busqueda}$2:${columna_busqueda}${ultima_fila},"
                f"{celda_actual})"
            )
            for columna_busqueda in columnas_ids
        ]
        condiciones.append(
            f'AND({celda_actual}<>"",SUM({",".join(conteos)})>1)'
        )

    return f'IF(OR({",".join(condiciones)}),"{mensaje_error}",'


def agregar_formulas_validacion_compuestos(ws_compuestos):
    if ws_compuestos.max_row < 2:
        return

    col_diferencia = get_column_letter(COLUMNAS_COMPUESTOS.index("Diferencia") + 1)
    col_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    col_validacion = get_column_letter(COLUMNAS_COMPUESTOS.index("Validacion") + 1)
    columnas_banco = [
        get_column_letter(COLUMNAS_COMPUESTOS.index(nombre) + 1)
        for nombre in ["id banco 1", "id banco 2", "id banco 3"]
    ]
    columnas_libro_aux = [
        get_column_letter(COLUMNAS_COMPUESTOS.index(nombre) + 1)
        for nombre in ["id libro_aux 1", "id libro_aux 2", "id libro_aux 3"]
    ]
    ultima_fila = ws_compuestos.max_row

    for fila in range(2, ultima_fila + 1):
        error_banco = crear_formula_id_repetido_en_compuestos(
            fila=fila,
            col_estado=col_estado,
            columnas_ids=columnas_banco,
            mensaje_error="ERROR: banco ya usado en otro compuesto conciliado",
            ultima_fila=ultima_fila,
        )
        error_libro_aux = crear_formula_id_repetido_en_compuestos(
            fila=fila,
            col_estado=col_estado,
            columnas_ids=columnas_libro_aux,
            mensaje_error="ERROR: libro_aux ya usado en otro compuesto conciliado",
            ultima_fila=ultima_fila,
        )

        ws_compuestos[f"{col_validacion}{fila}"] = (
            f'=IF({col_estado}{fila}<>"conciliado","Pendiente",'
            f'IF({col_diferencia}{fila}<>0,'
            f'"ERROR: diferencia distinta de cero",'
            f'{error_banco}{error_libro_aux}"OK"))))'
        )


def crear_formula_libro_aux_usado_en_compuesto(celda_id_libro_aux, ws_compuestos):
    col_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    col_validacion = get_column_letter(COLUMNAS_COMPUESTOS.index("Validacion") + 1)
    columnas_id_libro_aux = [
        get_column_letter(COLUMNAS_COMPUESTOS.index(nombre) + 1)
        for nombre in ["id libro_aux 1", "id libro_aux 2", "id libro_aux 3"]
    ]
    ultima_fila = max(ws_compuestos.max_row, 2)

    conteos = [
        (
            f"COUNTIFS('{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${col_estado}$2:${col_estado}${ultima_fila},"
            f'"conciliado",'
            f"'{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${col_validacion}$2:${col_validacion}${ultima_fila},"
            f'"OK",'
            f"'{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${columna_id}$2:${columna_id}${ultima_fila},"
            f"{celda_id_libro_aux})"
        )
        for columna_id in columnas_id_libro_aux
    ]

    return "SUM(" + ",".join(conteos) + ")>0"


def crear_formula_movimientos_banco_compuesto(celda_id_libro_aux, ws_compuestos):
    col_movimientos_banco = get_column_letter(
        COLUMNAS_COMPUESTOS.index("Movimientos banco") + 1
    )
    col_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    col_validacion = get_column_letter(COLUMNAS_COMPUESTOS.index("Validacion") + 1)
    columnas_id_libro_aux = [
        get_column_letter(COLUMNAS_COMPUESTOS.index(nombre) + 1)
        for nombre in ["id libro_aux 1", "id libro_aux 2", "id libro_aux 3"]
    ]
    ultima_fila = max(ws_compuestos.max_row, 2)
    condiciones_id = "+".join(
        [
            (
                f"('{HOJA_CONCILIACION_COMPUESTOS}'!"
                f"${columna_id}$2:${columna_id}${ultima_fila}={celda_id_libro_aux})"
            )
            for columna_id in columnas_id_libro_aux
        ]
    )

    return (
        f'TEXTJOIN(CHAR(10),TRUE,FILTER('
        f"'{HOJA_CONCILIACION_COMPUESTOS}'!"
        f"${col_movimientos_banco}$2:${col_movimientos_banco}${ultima_fila},"
        f"('{HOJA_CONCILIACION_COMPUESTOS}'!"
        f"${col_estado}$2:${col_estado}${ultima_fila}=\"conciliado\")*"
        f"('{HOJA_CONCILIACION_COMPUESTOS}'!"
        f"${col_validacion}$2:${col_validacion}${ultima_fila}=\"OK\")*"
        f"(({condiciones_id})>0)))"
    )


def crear_formula_banco_usado_en_compuesto(celda_id_banco, ws_compuestos):
    col_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    col_validacion = get_column_letter(COLUMNAS_COMPUESTOS.index("Validacion") + 1)
    columnas_id_banco = [
        get_column_letter(COLUMNAS_COMPUESTOS.index(nombre) + 1)
        for nombre in ["id banco 1", "id banco 2", "id banco 3"]
    ]
    ultima_fila = max(ws_compuestos.max_row, 2)

    conteos = [
        (
            f"COUNTIFS('{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${col_estado}$2:${col_estado}${ultima_fila},"
            f'"conciliado",'
            f"'{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${col_validacion}$2:${col_validacion}${ultima_fila},"
            f'"OK",'
            f"'{HOJA_CONCILIACION_COMPUESTOS}'!"
            f"${columna_id}$2:${columna_id}${ultima_fila},"
            f"{celda_id_banco})"
        )
        for columna_id in columnas_id_banco
    ]

    return "SUM(" + ",".join(conteos) + ")>0"


def agregar_formulas_registros_pendientes_banco(ws_registros, ws_compuestos):
    if ws_registros.max_row < FILA_DATOS_REGISTROS_BANCO:
        return

    for fila in range(FILA_DATOS_REGISTROS_BANCO, ws_registros.max_row + 1):
        celda_id_banco = f"A{fila}"
        banco_usado_en_compuesto = crear_formula_banco_usado_en_compuesto(
            celda_id_banco,
            ws_compuestos,
        )
        ws_registros[f"E{fila}"] = (
            f'=IF({banco_usado_en_compuesto},'
            f'"conciliado en compuesto confirmado",'
            f'"disponible")'
        )


def agregar_formulas_libro_auxiliar(ws_libro, ws_compuestos):
    if ws_libro.max_row < FILA_DATOS_LIBRO_AUXILIAR:
        return

    col_id = get_column_letter(1)
    col_estado = get_column_letter(COL_ESTADO_AUTOMATICO)
    col_conciliado_con = get_column_letter(COL_CONCILIADO_AUTOMATICO)
    col_estado_manual = get_column_letter(COL_ESTADO_MANUAL)
    col_estado_automatico_base = get_column_letter(COL_ESTADO_AUTOMATICO_BASE)

    for fila in range(FILA_DATOS_LIBRO_AUXILIAR, ws_libro.max_row + 1):
        celda_estado = f"{col_estado}{fila}"
        celda_conciliado_con = f"{col_conciliado_con}{fila}"
        celda_estado_manual = f"{col_estado_manual}{fila}"
        celda_estado_automatico_base = f"{col_estado_automatico_base}{fila}"

        if ws_libro[celda_estado].value == "conciliado":
            ws_libro[celda_estado] = (
                f'=IF({celda_estado_manual}="conciliado",'
                f'"conciliado",{celda_estado_automatico_base})'
            )
            ws_libro[celda_estado_manual] = (
                f'=IF({celda_estado_automatico_base}="conciliado",'
                f'"conciliado","por revisar")'
            )
            continue

        celda_id = f"{col_id}{fila}"
        compuesto_confirmado = crear_formula_libro_aux_usado_en_compuesto(
            celda_id,
            ws_compuestos,
        )
        movimientos_banco = crear_formula_movimientos_banco_compuesto(
            celda_id,
            ws_compuestos,
        )

        ws_libro[celda_estado_automatico_base] = (
            f'=IF({compuesto_confirmado},"conciliado","por revisar")'
        )
        ws_libro[celda_estado] = (
            f'=IF({celda_estado_manual}="conciliado",'
            f'"conciliado",'
            f'{celda_estado_automatico_base})'
        )
        ws_libro[celda_estado_manual] = (
            f'=IF({celda_estado_automatico_base}="conciliado",'
            f'"conciliado","por revisar")'
        )
        ws_libro[celda_conciliado_con] = (
            f'=IF({compuesto_confirmado},IFERROR({movimientos_banco},""),"")'
        )


def colorear_estado_libro_auxiliar(ws):
    if ws.max_row < FILA_DATOS_LIBRO_AUXILIAR:
        return

    columna_estado = get_column_letter(COL_ESTADO_AUTOMATICO)
    rango = (
        f"A{FILA_DATOS_LIBRO_AUXILIAR}:"
        f"{get_column_letter(COL_CONCILIADO_AUTOMATICO)}{ws.max_row}"
    )

    ws.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'${columna_estado}{FILA_DATOS_LIBRO_AUXILIAR}="conciliado"'],
            fill=PatternFill(
                start_color=COLOR_CONCILIADO,
                end_color=COLOR_CONCILIADO,
                fill_type="solid",
            ),
        ),
    )
    ws.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'${columna_estado}{FILA_DATOS_LIBRO_AUXILIAR}="por revisar"'],
            fill=PatternFill(
                start_color=COLOR_POR_REVISAR,
                end_color=COLOR_POR_REVISAR,
                fill_type="solid",
            ),
        ),
    )


def aplicar_estilo_libro_auxiliar(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{FILA_DATOS_LIBRO_AUXILIAR}"
    ws.auto_filter.ref = (
        f"A{FILA_ENCABEZADOS_LIBRO_AUXILIAR}:"
        f"{get_column_letter(COL_CONCILIADO_MANUAL)}{ws.max_row}"
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = "Conciliacion Bancaria"
    ws["A1"].fill = PatternFill(
        start_color=COLOR_TITULO,
        end_color=COLOR_TITULO,
        fill_type="solid",
    )
    ws["A1"].font = Font(
        bold=True,
        italic=True,
        size=TAMANO_FUENTE_TITULO,
        color="FFFFFF",
    )
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    bloques = [
        ("A2:D2", "Libro Auxiliar", COLOR_BLOQUE_LIBRO),
        ("E2:F2", "Conciliacion Automatica", COLOR_BLOQUE_AUTOMATICO),
        ("G2:H2", "Conciliacion manual", COLOR_BLOQUE_MANUAL),
    ]
    for rango, titulo, color in bloques:
        ws.merge_cells(rango)
        celda = ws[rango.split(":")[0]]
        celda.value = titulo
        celda.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        celda.font = Font(
            bold=True,
            size=TAMANO_FUENTE_ENCABEZADO,
            color="FFFFFF",
        )
        celda.alignment = Alignment(horizontal="center", vertical="center")

    borde_fino = Side(style="thin", color=COLOR_BORDE)
    borde = Border(
        left=borde_fino,
        right=borde_fino,
        top=borde_fino,
        bottom=borde_fino,
    )
    encabezado_fill = PatternFill(
        start_color=COLOR_ENCABEZADO,
        end_color=COLOR_ENCABEZADO,
        fill_type="solid",
    )

    for fila in range(1, FILA_ENCABEZADOS_LIBRO_AUXILIAR + 1):
        for celda in ws[fila]:
            celda.border = borde

    for celda in ws[FILA_ENCABEZADOS_LIBRO_AUXILIAR]:
        celda.fill = encabezado_fill
        celda.font = Font(
            bold=True,
            size=TAMANO_FUENTE_ENCABEZADO,
            color="FFFFFF",
        )
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    for fila in ws.iter_rows(min_row=FILA_DATOS_LIBRO_AUXILIAR):
        for celda in fila:
            celda.font = Font(size=TAMANO_FUENTE_DATOS)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    ws.cell(
        row=FILA_ENCABEZADOS_LIBRO_AUXILIAR,
        column=COL_CONCILIADO_MANUAL,
    ).comment = Comment(
        "Podrías usar un formato así:\n"
        "ID / dd-mm-yyyy / nombre / valor\n"
        "ID / dd-mm-yyyy / nombre / valor",
        "Conciliacion",
    )

    ws.column_dimensions[get_column_letter(COL_ESTADO_AUTOMATICO_BASE)].hidden = True

    aplicar_formatos_numericos(
        ws,
        fila_encabezados=FILA_ENCABEZADOS_LIBRO_AUXILIAR,
    )
    aplicar_anchos_libro_auxiliar(ws)


def aplicar_anchos_libro_auxiliar(ws):
    for columna, ancho in ANCHOS_LIBRO_AUXILIAR.items():
        ws.column_dimensions[columna].width = ancho


def agregar_validacion_estado_manual(ws):
    if ws.max_row < FILA_DATOS_LIBRO_AUXILIAR:
        return

    columna_estado = get_column_letter(COL_ESTADO_MANUAL)
    rango_destino = (
        f"{columna_estado}{FILA_DATOS_LIBRO_AUXILIAR}:"
        f"{columna_estado}{ws.max_row}"
    )
    validacion = DataValidation(
        type="list",
        formula1='"por revisar,conciliado"',
        allow_blank=True,
    )
    validacion.error = "Seleccioná conciliado o por revisar."
    validacion.errorTitle = "Estado inválido"
    validacion.prompt = "Usá conciliado solo cuando completes la revisión manual."
    validacion.promptTitle = "Estado manual"

    ws.add_data_validation(validacion)
    validacion.add(rango_destino)


def colorear_estado_manual(ws):
    if ws.max_row < FILA_DATOS_LIBRO_AUXILIAR:
        return

    columna_estado = get_column_letter(COL_ESTADO_MANUAL)
    rango_fila_completa = (
        f"A{FILA_DATOS_LIBRO_AUXILIAR}:"
        f"{get_column_letter(COL_CONCILIADO_MANUAL)}{ws.max_row}"
    )
    rango_manual = (
        f"{get_column_letter(COL_ESTADO_MANUAL)}{FILA_DATOS_LIBRO_AUXILIAR}:"
        f"{get_column_letter(COL_CONCILIADO_MANUAL)}{ws.max_row}"
    )

    ws.conditional_formatting.add(
        rango_fila_completa,
        FormulaRule(
            formula=[f'${columna_estado}{FILA_DATOS_LIBRO_AUXILIAR}="conciliado"'],
            fill=PatternFill(
                start_color=COLOR_CONCILIADO,
                end_color=COLOR_CONCILIADO,
                fill_type="solid",
            ),
        ),
    )
    ws.conditional_formatting.add(
        rango_manual,
        FormulaRule(
            formula=[f'${columna_estado}{FILA_DATOS_LIBRO_AUXILIAR}="por revisar"'],
            fill=PatternFill(
                start_color=COLOR_POR_REVISAR,
                end_color=COLOR_POR_REVISAR,
                fill_type="solid",
            ),
        ),
    )


def aplicar_estilo_registros_pendientes_banco(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{FILA_DATOS_REGISTROS_BANCO}"
    ws.auto_filter.ref = f"A{FILA_ENCABEZADOS_REGISTROS_BANCO}:E{ws.max_row}"

    ws.merge_cells("A1:E1")
    ws["A1"] = "Extractos bancarios"
    ws["A1"].fill = PatternFill(
        start_color=COLOR_BLOQUE_LIBRO,
        end_color=COLOR_BLOQUE_LIBRO,
        fill_type="solid",
    )
    ws["A1"].font = Font(
        bold=True,
        size=TAMANO_FUENTE_ENCABEZADO,
        color="FFFFFF",
    )
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    borde_fino = Side(style="thin", color=COLOR_BORDE)
    borde = Border(
        left=borde_fino,
        right=borde_fino,
        top=borde_fino,
        bottom=borde_fino,
    )
    encabezado_fill = PatternFill(
        start_color=COLOR_ENCABEZADO,
        end_color=COLOR_ENCABEZADO,
        fill_type="solid",
    )

    for fila in ws.iter_rows():
        for celda in fila:
            if celda.row >= FILA_DATOS_REGISTROS_BANCO:
                celda.font = Font(size=TAMANO_FUENTE_DATOS)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    for celda in ws[FILA_ENCABEZADOS_REGISTROS_BANCO]:
        celda.fill = encabezado_fill
        celda.font = Font(
            bold=True,
            size=TAMANO_FUENTE_ENCABEZADO,
            color="FFFFFF",
        )
        celda.alignment = Alignment(horizontal="center", vertical="center")

    aplicar_formatos_numericos(
        ws,
        fila_encabezados=FILA_ENCABEZADOS_REGISTROS_BANCO,
    )
    aplicar_anchos_registros_pendientes_banco(ws)


def aplicar_anchos_registros_pendientes_banco(ws):
    for columna, ancho in ANCHOS_REGISTROS_PENDIENTES_BANCO.items():
        ws.column_dimensions[columna].width = ancho


def agregar_validacion_estado_registros_pendientes_banco(ws):
    if ws.max_row < FILA_DATOS_REGISTROS_BANCO:
        return

    rango_destino = f"E{FILA_DATOS_REGISTROS_BANCO}:E{ws.max_row}"
    validacion = DataValidation(
        type="list",
        formula1=(
            '"disponible,conciliado manualmente,'
            'conciliado en compuesto confirmado"'
        ),
        allow_blank=True,
    )
    validacion.error = "Seleccioná un estado válido."
    validacion.errorTitle = "Estado inválido"
    validacion.prompt = "Usá conciliado manualmente cuando cierres el registro."
    validacion.promptTitle = "Estado banco"

    ws.add_data_validation(validacion)
    validacion.add(rango_destino)


def colorear_registros_pendientes_banco(ws):
    if ws.max_row < FILA_DATOS_REGISTROS_BANCO:
        return

    rango = f"A{FILA_DATOS_REGISTROS_BANCO}:E{ws.max_row}"

    ws.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'$E{FILA_DATOS_REGISTROS_BANCO}="disponible"'],
            fill=PatternFill(
                start_color=COLOR_POR_REVISAR,
                end_color=COLOR_POR_REVISAR,
                fill_type="solid",
            ),
        ),
    )
    ws.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[
                (
                    f'OR($E{FILA_DATOS_REGISTROS_BANCO}='
                    f'"conciliado manualmente",'
                    f'$E{FILA_DATOS_REGISTROS_BANCO}='
                    f'"conciliado en compuesto confirmado")'
                )
            ],
            fill=PatternFill(
                start_color=COLOR_CONCILIADO,
                end_color=COLOR_CONCILIADO,
                fill_type="solid",
            ),
        ),
    )


def colorear_conciliacion_compuestos(ws_compuestos):
    if ws_compuestos.max_row < 2:
        return

    col_estado = get_column_letter(COLUMNAS_COMPUESTOS.index("Estado") + 1)
    col_validacion = get_column_letter(COLUMNAS_COMPUESTOS.index("Validacion") + 1)
    ultima_columna_visible = get_column_letter(len(COLUMNAS_COMPUESTOS_VISIBLES))
    rango = f"A2:{ultima_columna_visible}{ws_compuestos.max_row}"

    ws_compuestos.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'LEFT(${col_validacion}2,5)="ERROR"'],
            fill=PatternFill(
                start_color=COLOR_ERROR,
                end_color=COLOR_ERROR,
                fill_type="solid",
            ),
        ),
    )
    ws_compuestos.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'AND(${col_estado}2="conciliado",${col_validacion}2="OK")'],
            fill=PatternFill(
                start_color=COLOR_OK,
                end_color=COLOR_OK,
                fill_type="solid",
            ),
        ),
    )
    ws_compuestos.conditional_formatting.add(
        rango,
        FormulaRule(
            formula=[f'${col_estado}2="por revisar"'],
            fill=PatternFill(
                start_color=COLOR_POR_REVISAR,
                end_color=COLOR_POR_REVISAR,
                fill_type="solid",
            ),
        ),
    )
    ws_compuestos.conditional_formatting.add(
        f"{col_validacion}2:{col_validacion}{ws_compuestos.max_row}",
        FormulaRule(
            formula=[f'${col_validacion}2="Pendiente"'],
            fill=PatternFill(
                start_color=COLOR_PENDIENTE,
                end_color=COLOR_PENDIENTE,
                fill_type="solid",
            ),
        ),
    )


def exportar_resultado_conciliacion(
    df_resultado,
    df_compuestos=None,
    nombre_archivo=None,
    df_conceptos_banco=None,
    df_banco_extraido=None,
    df_libro_auxiliar_extraido=None,
):
    """
    Exporta el resultado de conciliación a un Excel simple, auditable y claro.
    """

    if nombre_archivo is None and isinstance(df_compuestos, (str, Path)):
        nombre_archivo = df_compuestos
        df_compuestos = None

    if nombre_archivo is None:
        nombre_archivo = Path(__file__).with_name("resultado_conciliacion.xlsx")

    df_libro_auxiliar = preparar_libro_auxiliar_vs_banco(df_resultado)
    df_conceptos = preparar_conceptos_bancarios(df_conceptos_banco)
    df_registros_pendientes_banco = preparar_registros_pendientes_banco(df_resultado)
    df_compuestos_excel = preparar_conciliacion_compuestos(
        df_compuestos,
        df_resultado,
    )

    with pd.ExcelWriter(nombre_archivo, engine="openpyxl") as writer:
        ws_libro = escribir_dataframe(
            writer,
            df_libro_auxiliar,
            HOJA_LIBRO_AUXILIAR_BANCO,
            startrow=FILA_ENCABEZADOS_LIBRO_AUXILIAR - 1,
        )
        ws_compuestos = escribir_dataframe(
            writer,
            df_compuestos_excel,
            HOJA_CONCILIACION_COMPUESTOS,
        )
        ws_registros_banco = escribir_dataframe(
            writer,
            df_registros_pendientes_banco,
            HOJA_REGISTROS_PENDIENTES_BANCO,
            startrow=FILA_ENCABEZADOS_REGISTROS_BANCO - 1,
        )
        ws_conceptos = escribir_dataframe(
            writer,
            df_conceptos,
            HOJA_CONCEPTOS_BANCARIOS,
        )
        ws_banco_extraido = escribir_documento_extraido(
            writer,
            df_banco_extraido,
            HOJA_EXTRACTOS_BANCO,
        )
        ws_libro_auxiliar_extraido = escribir_documento_extraido(
            writer,
            df_libro_auxiliar_extraido,
            HOJA_LIBRO_AUXILIAR_EXTRAIDO,
        )

        aplicar_estilo_libro_auxiliar(ws_libro)
        agregar_formulas_libro_auxiliar(ws_libro, ws_compuestos)
        colorear_estado_libro_auxiliar(ws_libro)
        agregar_validacion_estado_manual(ws_libro)
        colorear_estado_manual(ws_libro)
        aplicar_anchos_libro_auxiliar(ws_libro)

        aplicar_estilo_base(ws_conceptos)
        ajustar_ancho_columnas(ws_conceptos)

        if ws_banco_extraido is not None:
            aplicar_estilo_base(ws_banco_extraido)

        if ws_libro_auxiliar_extraido is not None:
            aplicar_estilo_base(ws_libro_auxiliar_extraido)

        aplicar_estilo_base(ws_compuestos)
        agregar_validacion_estado_compuesto(ws_compuestos)
        agregar_formulas_validacion_compuestos(ws_compuestos)
        colorear_conciliacion_compuestos(ws_compuestos)
        ocultar_columnas(ws_compuestos, COLUMNAS_COMPUESTOS_AUXILIARES)
        aplicar_anchos_conciliacion_compuestos(ws_compuestos)

        aplicar_estilo_registros_pendientes_banco(ws_registros_banco)
        agregar_formulas_registros_pendientes_banco(
            ws_registros_banco,
            ws_compuestos,
        )
        agregar_validacion_estado_registros_pendientes_banco(ws_registros_banco)
        colorear_registros_pendientes_banco(ws_registros_banco)

    print(f"Archivo exportado: {nombre_archivo}")
