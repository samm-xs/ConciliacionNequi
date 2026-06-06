from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from Backend.ExportacionExcel.Exportar_resultado_conciliacion import (
    ANCHOS_CONCILIACION_COMPUESTOS,
    ANCHOS_LIBRO_AUXILIAR,
    ANCHOS_REGISTROS_PENDIENTES_BANCO,
    HOJA_CONCEPTOS_BANCARIOS,
    HOJA_CONCILIACION_COMPUESTOS,
    HOJA_EXTRACTOS_BANCO,
    HOJA_LIBRO_AUXILIAR_EXTRAIDO,
    HOJA_LIBRO_AUXILIAR_BANCO,
    HOJA_REGISTROS_PENDIENTES_BANCO,
    TAMANO_FUENTE_DATOS,
    TAMANO_FUENTE_ENCABEZADO,
    TAMANO_FUENTE_TITULO,
    exportar_resultado_conciliacion,
)


def construir_resultado_base():
    return pd.DataFrame(
        [
            {
                "id_banco": 1,
                "fecha_banco": "2026-03-01",
                "nombre_banco": "cliente uno",
                "valor_banco": 50000,
                "id_erp": 10,
                "fecha_erp": "2026-03-01",
                "nombre_erp": "cliente uno",
                "valor_erp": 50000,
                "estado": "conciliado",
                "observacion": "Conciliación automática.",
            },
            {
                "id_banco": 2,
                "fecha_banco": "2026-03-02",
                "nombre_banco": "pago parcial a",
                "valor_banco": 20000,
                "id_erp": None,
                "fecha_erp": None,
                "nombre_erp": None,
                "valor_erp": None,
                "estado": "huerfano_banco",
                "observacion": "Banco sin candidato.",
            },
            {
                "id_banco": 3,
                "fecha_banco": "2026-03-02",
                "nombre_banco": "pago parcial b",
                "valor_banco": 30000,
                "id_erp": None,
                "fecha_erp": None,
                "nombre_erp": None,
                "valor_erp": None,
                "estado": "huerfano_banco",
                "observacion": "Banco sin candidato.",
            },
            {
                "id_banco": None,
                "fecha_banco": None,
                "nombre_banco": None,
                "valor_banco": None,
                "id_erp": 20,
                "fecha_erp": "2026-03-02",
                "nombre_erp": "proveedor compuesto",
                "valor_erp": 50000,
                "estado": "huerfano_erp",
                "observacion": "ERP pendiente.",
            },
            {
                "id_banco": None,
                "fecha_banco": None,
                "nombre_banco": None,
                "valor_banco": None,
                "id_erp": 30,
                "fecha_erp": "2026-03-03",
                "nombre_erp": "proveedor pendiente",
                "valor_erp": 70000,
                "estado": "huerfano_erp",
                "observacion": "ERP pendiente.",
            },
        ]
    )


def construir_compuestos_base():
    return pd.DataFrame(
        [
            {
                "tipo": "varios_banco_a_un_erp",
                "ids_banco": "2, 3",
                "ids_erp": "20",
                "fecha_referencia": "2026-03-02",
                "valor_total_banco": Decimal("50000"),
                "valor_total_erp": Decimal("50000"),
                "diferencia": Decimal("0"),
                "cantidad_banco": 2,
                "cantidad_erp": 1,
                "estado_compuesto": "Por revisar",
                "validacion_compuesto": "OK",
                "descripcion": (
                    "Posible compuesto: varios movimientos banco "
                    "suman un ERP. Revisar antes de conciliar."
                ),
            }
        ]
    )


def construir_conceptos_banco_base():
    return pd.DataFrame(
        [
            {
                "concepto": "al gravamen movimiento",
                "cantidad movimientos": 2,
                "valor total": -1200,
            },
            {
                "concepto": "de intereses pago",
                "cantidad movimientos": 1,
                "valor total": 25,
            },
        ]
    )


def encabezados(ws, fila=1):
    return [celda.value for celda in ws[fila]]


def rangos_validacion(ws):
    return [
        str(rango)
        for validacion in ws.data_validations.dataValidation
        for rango in validacion.cells.ranges
    ]


def exportar_base(tmp_path, conceptos=None):
    archivo = tmp_path / "resultado.xlsx"
    exportar_resultado_conciliacion(
        construir_resultado_base(),
        construir_compuestos_base(),
        nombre_archivo=archivo,
        df_conceptos_banco=conceptos,
    )
    return archivo


def test_exporta_workbook_simplificado_con_hoja_de_bancos_pendientes(tmp_path):
    archivo = exportar_base(tmp_path)

    wb = load_workbook(archivo, data_only=False)

    assert wb.sheetnames == [
        HOJA_LIBRO_AUXILIAR_BANCO,
        HOJA_CONCILIACION_COMPUESTOS,
        HOJA_REGISTROS_PENDIENTES_BANCO,
        HOJA_CONCEPTOS_BANCARIOS,
    ]
    assert "Banco vs ERP" not in wb.sheetnames
    assert "Conciliación manual" not in wb.sheetnames
    assert "ERP pendientes" not in wb.sheetnames
    assert "ERP disponibles" not in wb.sheetnames
    assert "Posibles compuestos" not in wb.sheetnames


def test_exporta_hojas_de_documentos_extraidos_crudos(tmp_path):
    archivo = tmp_path / "resultado.xlsx"
    df_banco_extraido = pd.DataFrame(
        [
            {
                "fecha original": "2026-03-01",
                "descripcion original": "Pago proveedor",
                "valor original": 50000,
            }
        ]
    )
    df_libro_auxiliar_extraido = pd.DataFrame(
        [
            {
                "fecha factura": "2026-03-01",
                "tercero": "Proveedor SAS",
                "total": 50000,
            }
        ]
    )

    exportar_resultado_conciliacion(
        construir_resultado_base(),
        construir_compuestos_base(),
        nombre_archivo=archivo,
        df_banco_extraido=df_banco_extraido,
        df_libro_auxiliar_extraido=df_libro_auxiliar_extraido,
    )

    wb = load_workbook(archivo, data_only=False)
    ws_banco = wb[HOJA_EXTRACTOS_BANCO]
    ws_libro = wb[HOJA_LIBRO_AUXILIAR_EXTRAIDO]

    assert HOJA_LIBRO_AUXILIAR_BANCO in wb.sheetnames
    assert HOJA_EXTRACTOS_BANCO in wb.sheetnames
    assert HOJA_LIBRO_AUXILIAR_EXTRAIDO in wb.sheetnames
    assert encabezados(ws_banco) == [
        "fecha original",
        "descripcion original",
        "valor original",
    ]
    assert ws_banco["A2"].value == "2026-03-01"
    assert ws_banco["B2"].value == "Pago proveedor"
    assert ws_banco["C2"].value == 50000
    assert encabezados(ws_libro) == ["fecha factura", "tercero", "total"]
    assert ws_libro["A2"].value == "2026-03-01"
    assert ws_libro["B2"].value == "Proveedor SAS"
    assert ws_libro["C2"].value == 50000
    assert ws_banco.freeze_panes == "A2"
    assert ws_libro.freeze_panes == "A2"
    assert ws_banco.auto_filter.ref == "A1:C2"
    assert ws_libro.auto_filter.ref == "A1:C2"
    assert ws_banco["A1"].font.sz == TAMANO_FUENTE_ENCABEZADO
    assert ws_banco["A2"].font.sz == TAMANO_FUENTE_DATOS
    assert ws_libro["A1"].font.sz == TAMANO_FUENTE_ENCABEZADO
    assert ws_libro["A2"].font.sz == TAMANO_FUENTE_DATOS


def test_libro_auxiliar_tiene_columnas_y_una_fila_por_id_erp(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["A1"].value == "Conciliacion Bancaria"
    assert ws["A2"].value == "Libro Auxiliar"
    assert ws["E2"].value == "Conciliacion Automatica"
    assert ws["G2"].value == "Conciliacion manual"
    assert encabezados(ws, fila=3)[:8] == [
        "ID",
        "Fecha",
        "Nombre",
        "Valor",
        "Estado",
        "Conciliado del banco con registros:",
        "Estado",
        "Conciliado del banco con registros:",
    ]
    assert ws["I3"].value == "Estado automatico base"
    assert ws.column_dimensions["I"].hidden
    assert ws.max_row == 6
    assert [ws[f"A{fila}"].value for fila in range(4, 7)] == [10, 20, 30]


def test_libro_auxiliar_tiene_titulo_y_anchos_personalizados(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["A1"].fill.fgColor.rgb == "00A64D79"
    assert ws["A1"].font.color.rgb == "00FFFFFF"
    assert ws["A1"].font.bold
    assert ws["A1"].font.italic
    assert ws["A1"].font.sz == TAMANO_FUENTE_TITULO
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A1"].alignment.vertical == "center"
    for columna, ancho in ANCHOS_LIBRO_AUXILIAR.items():
        assert ws.column_dimensions[columna].width == ancho


def test_libro_auxiliar_aplica_tamanos_de_fuente_estandarizados(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["A1"].font.sz == TAMANO_FUENTE_TITULO
    assert ws["A2"].font.sz == TAMANO_FUENTE_ENCABEZADO
    assert ws["A3"].font.sz == TAMANO_FUENTE_ENCABEZADO
    assert ws["A4"].font.sz == TAMANO_FUENTE_DATOS


def test_libro_auxiliar_muestra_conciliacion_automatica_uno_a_uno(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["E4"].value.startswith("=IF(")
    assert 'G4="conciliado"' in ws["E4"].value
    assert "I4" in ws["E4"].value
    assert ws["F4"].value == "1 | 2026-03-01 | 50,000 | cliente uno"
    assert ws["G4"].value == '=IF(I4="conciliado","conciliado","por revisar")'
    assert ws["I4"].value == "conciliado"


def test_libro_auxiliar_pendiente_refleja_compuesto_confirmado_por_formula(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["E5"].value.startswith("=IF(")
    assert 'G5="conciliado"' in ws["E5"].value
    assert "I5" in ws["E5"].value
    assert ws["I5"].value.startswith("=IF(")
    assert "Conciliacion de compuestos" in ws["I5"].value
    assert '"conciliado"' in ws["E5"].value
    assert '"por revisar"' in ws["I5"].value
    assert ws["G5"].value == '=IF(I5="conciliado","conciliado","por revisar")'
    assert ws["F5"].value.startswith("=IF(")
    assert "TEXTJOIN" in ws["F5"].value
    assert "FILTER" in ws["F5"].value
    assert "Conciliacion de compuestos" in ws["F5"].value


def test_libro_auxiliar_pendiente_sin_compuesto_queda_por_revisar_por_formula(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]

    assert ws["A6"].value == 30
    assert ws["E6"].value.startswith("=IF(")
    assert 'G6="conciliado"' in ws["E6"].value
    assert "I6" in ws["E6"].value
    assert ws["G6"].value == '=IF(I6="conciliado","conciliado","por revisar")'
    assert '"por revisar"' in ws["I6"].value


def test_libro_auxiliar_tiene_dropdown_manual_y_guia_de_formato(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]
    validaciones = ws.data_validations.dataValidation

    assert any("G4:G6" in str(rango) for rango in rangos_validacion(ws))
    assert validaciones[0].formula1 == '"por revisar,conciliado"'
    assert ws["H4"].value is None
    assert ws["H3"].comment is not None
    assert "Podrías usar un formato así:" in ws["H3"].comment.text
    assert "ID / dd-mm-yyyy / nombre / valor" in ws["H3"].comment.text


def test_libro_auxiliar_formato_manual_pinta_fila_completa(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_LIBRO_AUXILIAR_BANCO]
    rangos = [str(rango) for rango in ws.conditional_formatting]

    assert any("A4:H6" in rango for rango in rangos)
    assert any("A4:F6" in rango for rango in rangos)


def test_conciliacion_compuestos_tiene_columnas_dropdown_y_formula(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCILIACION_COMPUESTOS]
    headers = encabezados(ws)
    col_validacion = headers.index("Validacion") + 1

    assert headers[:7] == [
        "Movimientos banco",
        "Movimientos libro_aux",
        "Diferencia",
        "Estado",
        "Validacion",
        "Descripcion",
        "Observacion",
    ]
    assert "D2" in rangos_validacion(ws)
    assert ws.data_validations.dataValidation[0].formula1 == '"por revisar,conciliado"'
    assert ws.cell(row=2, column=col_validacion).value.startswith("=IF(")
    assert "ERROR: diferencia distinta de cero" in ws.cell(
        row=2,
        column=col_validacion,
    ).value
    assert "ERROR: banco ya usado en otro compuesto conciliado" in ws.cell(
        row=2,
        column=col_validacion,
    ).value
    assert "ERROR: libro_aux ya usado en otro compuesto conciliado" in ws.cell(
        row=2,
        column=col_validacion,
    ).value


def test_conciliacion_compuestos_muestra_movimientos_compactos_y_oculta_ids(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCILIACION_COMPUESTOS]
    headers = encabezados(ws)

    assert "2 | 2026-03-02 | 20,000 | pago parcial a" in ws["A2"].value
    assert "3 | 2026-03-02 | 30,000 | pago parcial b" in ws["A2"].value
    assert "20 | 2026-03-02 | 50,000 | proveedor compuesto" == ws["B2"].value
    assert "ids banco" in headers
    assert "ids libro_aux" in headers
    assert ws.column_dimensions["H"].hidden
    assert ws.column_dimensions["I"].hidden
    assert ws.column_dimensions["J"].hidden
    assert ws.column_dimensions["O"].hidden


def test_conciliacion_compuestos_tiene_anchos_configurables(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCILIACION_COMPUESTOS]

    assert ws.column_dimensions["A"].width == ANCHOS_CONCILIACION_COMPUESTOS["A"]
    assert ws.column_dimensions["B"].width == ANCHOS_CONCILIACION_COMPUESTOS["B"]
    assert ws.column_dimensions["C"].width == ANCHOS_CONCILIACION_COMPUESTOS["C"]
    assert ws.column_dimensions["D"].width == ANCHOS_CONCILIACION_COMPUESTOS["D"]
    assert ws.column_dimensions["E"].width == ANCHOS_CONCILIACION_COMPUESTOS["E"]
    assert ws.column_dimensions["F"].width == ANCHOS_CONCILIACION_COMPUESTOS["F"]
    assert ws.column_dimensions["G"].width == ANCHOS_CONCILIACION_COMPUESTOS["G"]
    assert ws.column_dimensions["H"].width == ANCHOS_CONCILIACION_COMPUESTOS["H"]
    assert ws.column_dimensions["O"].width == ANCHOS_CONCILIACION_COMPUESTOS["O"]
    assert ws.column_dimensions["H"].hidden
    assert ws.column_dimensions["O"].hidden


def test_registros_pendientes_banco_muestra_solo_bancos_no_conciliados(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_REGISTROS_PENDIENTES_BANCO]

    assert ws["A1"].value == "Extractos bancarios"
    assert encabezados(ws, fila=2) == [
        "ID",
        "Fecha",
        "Nombre",
        "Valor",
        "Estado",
    ]
    assert [ws[f"A{fila}"].value for fila in range(3, 5)] == [2, 3]
    assert ws.max_row == 4
    assert 1 not in [ws[f"A{fila}"].value for fila in range(3, ws.max_row + 1)]


def test_registros_pendientes_banco_estado_es_editable_y_mira_compuestos(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_REGISTROS_PENDIENTES_BANCO]

    assert any("E3:E4" in str(rango) for rango in rangos_validacion(ws))
    assert (
        ws.data_validations.dataValidation[0].formula1
        == '"disponible,conciliado manualmente,conciliado en compuesto confirmado"'
    )
    assert ws["E3"].value.startswith("=IF(")
    assert "Conciliacion de compuestos" in ws["E3"].value
    assert '"conciliado"' in ws["E3"].value
    assert '"OK"' in ws["E3"].value
    assert '"conciliado en compuesto confirmado"' in ws["E3"].value
    assert '"disponible"' in ws["E3"].value
    assert "$J$2:$J$2" in ws["E3"].value
    assert "$K$2:$K$2" in ws["E3"].value
    assert "$L$2:$L$2" in ws["E3"].value


def test_registros_pendientes_banco_tiene_estilo_y_colores_por_estado(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_REGISTROS_PENDIENTES_BANCO]
    rangos = [str(rango) for rango in ws.conditional_formatting]

    assert ws.freeze_panes == "A3"
    assert ws.auto_filter.ref == "A2:E4"
    assert not ws.sheet_view.showGridLines
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A1"].alignment.vertical == "center"
    for columna, ancho in ANCHOS_REGISTROS_PENDIENTES_BANCO.items():
        assert ws.column_dimensions[columna].width == ancho
    assert ws["A3"].border.left.style == "thin"
    assert ws["A3"].border.left.color.rgb == "00000000"
    assert any("A3:E4" in rango for rango in rangos)


def test_limita_compuestos_exportados_para_google_sheets(tmp_path):
    archivo = tmp_path / "resultado.xlsx"
    compuestos = pd.concat(
        [construir_compuestos_base()] * 1001,
        ignore_index=True,
    )

    exportar_resultado_conciliacion(
        construir_resultado_base(),
        compuestos,
        nombre_archivo=archivo,
    )

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCILIACION_COMPUESTOS]

    assert ws.max_row == 1001


def test_exporta_hoja_conceptos_bancarios(tmp_path):
    archivo = exportar_base(tmp_path, conceptos=construir_conceptos_banco_base())

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCEPTOS_BANCARIOS]

    assert encabezados(ws) == [
        "Concepto",
        "Cantidad movimientos",
        "Valor total",
        "Descripcion",
    ]
    assert ws["A2"].value == "al gravamen movimiento"
    assert ws["B2"].value == 2
    assert ws["C2"].value == -1200
    assert (
        ws["D2"].value
        == "Concepto bancario informativo. No se concilia contra libro auxiliar."
    )
    assert ws["A1"].font.sz == TAMANO_FUENTE_ENCABEZADO
    assert ws["A2"].font.sz == TAMANO_FUENTE_DATOS


def test_exporta_hoja_conceptos_bancarios_vacia_si_no_se_entrega_dataframe(tmp_path):
    archivo = exportar_base(tmp_path)

    ws = load_workbook(archivo, data_only=False)[HOJA_CONCEPTOS_BANCARIOS]

    assert encabezados(ws) == [
        "Concepto",
        "Cantidad movimientos",
        "Valor total",
        "Descripcion",
    ]
    assert ws.max_row == 1
